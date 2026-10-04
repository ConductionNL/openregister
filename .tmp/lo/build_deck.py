import struct, zlib, io
from pptx import Presentation
from pptx.util import Inches

def png(w=2, h=2):
    raw = b''.join(b'\x00' + b'\x00\x80\x00' * w for _ in range(h))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
s1.shapes.title.text = 'Fotosynthese'
tf = s1.placeholders[1].text_frame
tf.text = 'Planten maken voedsel'
p = tf.add_paragraph(); p.text = 'Licht, water en CO2'
pic = s1.shapes.add_picture(io.BytesIO(png()), Inches(6), Inches(4), Inches(1), Inches(1))
pic.name = 'Blad'
pic._element.nvPicPr.cNvPr.set('descr', 'Een blad in de zon')
s1.notes_slide.notes_text_frame.text = 'Vraag eerst wat ze al weten.'
s1.notes_slide.notes_text_frame.add_paragraph().text = 'Laat ze daarna tekenen.'

s2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
s2.shapes.title.text = 'Water kookt'
table = s2.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
table.cell(0, 0).text = 'Cel A'; table.cell(0, 1).text = 'Cel B'

s3 = prs.slides.add_slide(prs.slide_layouts[1])
s3.shapes.title.text = 'Verborgen dia'
s3._element.set('show', '0')
prs.save('python-pptx.pptx')
print('ok')
